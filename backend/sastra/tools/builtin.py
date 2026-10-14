"""Built-in tools. Each handler receives (CallContext, validated args)."""
from __future__ import annotations

import csv
import io
import json
import statistics
import urllib.request
from typing import Literal

from pydantic import BaseModel, Field

from ..controls import Mode
from ..security.netguard import safe_fetch, validate_url
from ..security.untrusted import wrap_untrusted
from ..research.extract import extract
from ..workspace import android_template
from .calculator import calculate
from .registry import ToolError, ToolRegistry, ToolSpec


def _svc(ctx):
    return ctx.services


def _need_project(ctx) -> str:
    if not ctx.project_id:
        raise ToolError("This tool needs a project (create one in Projects and run the task inside it)")
    return ctx.project_id


# ------------------------------------------------------------------ inputs
class CalcIn(BaseModel):
    expression: str = Field(max_length=300)


class MemSearchIn(BaseModel):
    query: str = Field(max_length=500)
    limit: int = Field(6, ge=1, le=20)


class MemProposeIn(BaseModel):
    content: str = Field(max_length=4000)
    category: Literal["experimental", "working", "summary", "procedure"] = "experimental"
    subject: str = Field("", max_length=200)


class SearchIn(BaseModel):
    query: str = Field(max_length=400)
    count: int = Field(8, ge=1, le=20)
    freshness_days: int | None = Field(None, ge=1, le=3650)


class FetchIn(BaseModel):
    url: str = Field(max_length=2000)


class ResearchIn(BaseModel):
    question: str = Field(max_length=1000)
    urls: list[str] = Field(default_factory=list, max_length=10)
    max_sources: int = Field(6, ge=1, le=10)
    freshness_days: int | None = Field(None, ge=1, le=3650)


class FsListIn(BaseModel):
    pass


class FsReadIn(BaseModel):
    path: str = Field(max_length=500)


class FsWriteIn(BaseModel):
    path: str = Field(max_length=500)
    content: str = Field(max_length=2_000_000)


class CodeRunIn(BaseModel):
    language: Literal["python", "javascript", "bash"] = "python"
    path: str | None = Field(None, max_length=500, description="file in the project to run")
    code: str | None = Field(None, max_length=200_000, description="or inline code (saved to .sastra_run.*)")
    args: list[str] = Field(default_factory=list, max_length=20)
    timeout_s: int = Field(30, ge=1, le=600)
    network: Literal["none", "packages"] = "none"


class CheckpointIn(BaseModel):
    message: str = Field("Checkpoint", max_length=200)


class RollbackIn(BaseModel):
    commit: str = Field(pattern=r"^[0-9a-f]{7,40}$")


class SheetIn(BaseModel):
    path: str = Field(max_length=500)
    sheet: str | None = None
    max_rows: int = Field(100_000, ge=1, le=1_000_000)


class DocIn(BaseModel):
    path: str = Field(max_length=500)
    title: str = Field(max_length=200)
    markdown: str = Field(max_length=500_000)
    format: Literal["md", "html", "docx"] = "md"


class ArtifactIn(BaseModel):
    path: str = Field(max_length=500)
    name: str | None = Field(None, max_length=200)
    encrypt: bool = False


class AndroidIn(BaseModel):
    app_name: str = Field(max_length=40)
    package: str = Field(max_length=100)


class NotifyIn(BaseModel):
    title: str = Field(max_length=120)
    body: str = Field("", max_length=2000)


class ScheduleIn(BaseModel):
    name: str = Field(max_length=120)
    request: str = Field(max_length=2000)
    kind: Literal["research", "direct", "data", "document"] = "research"
    interval_minutes: int | None = Field(None, ge=60, le=60 * 24 * 31)
    daily_at: str | None = Field(None, pattern=r"^\d{2}:\d{2}$")
    budget_usd: float = Field(0.05, ge=0, le=5)


class WebhookIn(BaseModel):
    url: str = Field(max_length=2000)
    text: str = Field(max_length=20000)


class ProjectDeleteIn(BaseModel):
    project_id: str = Field(pattern=r"^[0-9a-f]{32}$")


# ---------------------------------------------------------------- handlers
def h_calc(ctx, a: CalcIn):
    return {"expression": a.expression, "result": calculate(a.expression)}


def h_mem_search(ctx, a: MemSearchIn):
    hits = _svc(ctx).memory.search(a.query, ctx.project_id, limit=a.limit)
    return {"results": [{"id": h.record["id"], "category": h.record["category"], "content": h.record["content"],
                         "score": h.score} for h in hits]}


def h_mem_propose(ctx, a: MemProposeIn):
    rec = _svc(ctx).memory.add(a.content, a.category, source="agent", subject=a.subject, project_id=ctx.project_id,
                               provenance={"task_id": ctx.task_id}, confidence=0.6, actor=ctx.actor)
    return {"id": rec["id"], "status": rec["status"]}


def h_search(ctx, a: SearchIn):
    results, notes = _svc(ctx).research.search(a.query, a.count, a.freshness_days, ctx.task_id)
    return {"results": [r.to_dict() for r in results], "notes": notes}


def h_fetch(ctx, a: FetchIn):
    s = _svc(ctx).settings
    fr = safe_fetch(a.url, max_bytes=s.fetch_max_bytes, timeout=s.fetch_timeout_s, respect_robots=s.respect_robots)
    ex = extract(fr.body, fr.content_type, fr.final_url)
    _svc(ctx).privacy.record(fr.final_url[:200], "web", "public", len(a.url), "fetch page", ctx.task_id)
    return {"url": fr.final_url, "status": fr.status, "title": ex.title, "published": ex.published,
            "content": wrap_untrusted(ex.text[:8000], fr.final_url), "truncated": fr.truncated or len(ex.text) > 8000}


def h_research(ctx, a: ResearchIn):
    out = _svc(ctx).research.run(a.question, urls=a.urls, max_sources=a.max_sources, freshness_days=a.freshness_days,
                                 project_id=ctx.project_id, task_id=ctx.task_id)
    return {"report_id": out["id"], "report_md": out["report_md"], "source_count": len(out["sources"]),
            "conflicts": len(out["conflicts"]), "synthesis": out["synthesis"], "notes": out["notes"]}


def h_fs_list(ctx, a: FsListIn):
    return {"files": _svc(ctx).projects.tree(_need_project(ctx))}


def h_fs_read(ctx, a: FsReadIn):
    pid = _need_project(ctx)
    return {"path": a.path, "content": _svc(ctx).projects.read(pid, a.path)}


def h_fs_write(ctx, a: FsWriteIn):
    pid = _need_project(ctx)
    p = _svc(ctx).projects
    p.checkpoint(pid, f"Before writing {a.path}", ctx.actor)
    out = p.write(pid, a.path, a.content, ctx.actor)
    return {"path": a.path, "diff": out["diff"][:20000]}


def h_code_run(ctx, a: CodeRunIn):
    svc = _svc(ctx)
    pid = _need_project(ctx)
    runner = svc.sandbox
    if runner is None:
        raise ToolError("No sandbox available: install Docker (recommended) or set SASTRA_ALLOW_WEAK_SANDBOX=1 for development only")
    if not (a.path or a.code):
        raise ToolError("Give a path or inline code")
    path = a.path
    if a.code:
        ext = {"python": "py", "javascript": "js", "bash": "sh"}[a.language]
        path = f".sastra_run.{ext}"
        svc.projects.write(pid, path, a.code, ctx.actor)
    svc.projects.resolve(pid, path)  # validates containment
    res = runner.run(svc.projects.dir(pid), a.language, [path, *a.args], a.timeout_s, a.network)
    svc.metrics.observe("sandbox.duration_ms", res.duration_ms, backend=res.backend)
    return res.to_dict()


def h_checkpoint(ctx, a: CheckpointIn):
    return {"commit": _svc(ctx).projects.checkpoint(_need_project(ctx), a.message, ctx.actor)}


def h_git_log(ctx, a: FsListIn):
    return {"log": _svc(ctx).projects.log(_need_project(ctx))}


def h_rollback(ctx, a: RollbackIn):
    return {"commit": _svc(ctx).projects.rollback(_need_project(ctx), a.commit, ctx.actor)}


def h_secrets(ctx, a: FsListIn):
    findings = _svc(ctx).projects.scan_secrets(_need_project(ctx))
    return {"findings": findings, "clean": not findings}


def h_sheet(ctx, a: SheetIn):
    pid = _need_project(ctx)
    p = _svc(ctx).projects.resolve(pid, a.path)
    rows: list[list] = []
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            import openpyxl
        except ImportError as e:
            raise ToolError("openpyxl not installed; save the sheet as CSV") from e
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        ws = wb[a.sheet] if a.sheet else wb.active
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if i > a.max_rows:
                break
            rows.append(list(r))
    else:
        text = p.read_text(encoding="utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|") if text.strip() else csv.excel
        for i, r in enumerate(csv.reader(io.StringIO(text), dialect)):
            if i > a.max_rows:
                break
            rows.append(r)
    if not rows:
        return {"rows": 0, "columns": []}
    header = [str(h) if h is not None else f"col{i + 1}" for i, h in enumerate(rows[0])]
    body = rows[1:]
    cols = []
    for ci, name in enumerate(header):
        vals = [r[ci] if ci < len(r) else None for r in body]
        missing = sum(1 for v in vals if v in (None, ""))
        nums = []
        for v in vals:
            try:
                nums.append(float(str(v).replace(",", "")))
            except (TypeError, ValueError):
                pass
        info = {"name": name, "missing": missing, "distinct": len({str(v) for v in vals})}
        if nums and len(nums) >= 0.8 * (len(vals) - missing or 1):
            info.update(type="number", min=min(nums), max=max(nums), mean=round(statistics.fmean(nums), 4),
                        sum=round(sum(nums), 4), median=statistics.median(nums))
        else:
            info["type"] = "text"
            info["top"] = sorted({str(v) for v in vals if v not in (None, "")})[:5]
        cols.append(info)
    dupes = len(body) - len({tuple(map(str, r)) for r in body})
    return {"rows": len(body), "columns": cols, "duplicate_rows": dupes, "sample": [dict(zip(header, map(str, r), strict=False)) for r in body[:5]]}


def h_doc(ctx, a: DocIn):
    pid = _need_project(ctx)
    projects = _svc(ctx).projects
    md = f"# {a.title}\n\n{a.markdown}" if not a.markdown.lstrip().startswith("#") else a.markdown
    if a.format == "md":
        projects.write(pid, a.path, md, ctx.actor)
    elif a.format == "html":
        import html as _h
        body = "".join(f"<h{ln.count('#', 0, 6)}>{_h.escape(ln.lstrip('# '))}</h{ln.count('#', 0, 6)}>" if ln.startswith("#")
                       else f"<p>{_h.escape(ln)}</p>" for ln in md.splitlines() if ln.strip())
        projects.write(pid, a.path, f"<!doctype html><meta charset=utf-8><title>{_h.escape(a.title)}</title>{body}", ctx.actor)
    else:
        try:
            import docx  # python-docx
        except ImportError as e:
            raise ToolError("python-docx not installed; use md or html") from e
        d = docx.Document()
        for ln in md.splitlines():
            if ln.startswith("#"):
                d.add_heading(ln.lstrip("# "), level=min(ln.count("#", 0, 6), 4))
            elif ln.strip().startswith(("- ", "* ")):
                d.add_paragraph(ln.strip()[2:], style="List Bullet")
            elif ln.strip():
                d.add_paragraph(ln)
        target = projects.resolve(pid, a.path)
        target.parent.mkdir(parents=True, exist_ok=True)
        d.save(target)
    return {"path": a.path, "format": a.format}


def h_artifact(ctx, a: ArtifactIn):
    pid = _need_project(ctx)
    return _svc(ctx).projects.store_artifact(pid, a.path, a.name or a.path.rsplit("/", 1)[-1], ctx.actor, ctx.task_id, a.encrypt)


def h_android(ctx, a: AndroidIn):
    pid = _need_project(ctx)
    p = _svc(ctx).projects
    written = []
    for rel, content in android_template.files(a.app_name, a.package).items():
        p.write(pid, rel, content, ctx.actor)
        written.append(rel)
    p.checkpoint(pid, f"Generated Android project {a.app_name}", ctx.actor)
    return {"files": written, "build": "Push to GitHub; the included workflow builds the APK (see BUILD_APK.md)"}


def h_notify(ctx, a: NotifyIn):
    nid = _svc(ctx).notifications.add("agent", a.title, a.body, ctx.task_id)
    return {"notification_id": nid}


def h_schedule(ctx, a: ScheduleIn):
    if not (a.interval_minutes or a.daily_at):
        raise ToolError("Give interval_minutes or daily_at")
    sid = _svc(ctx).scheduler.create(a.name, a.request, a.kind, ctx.project_id, a.interval_minutes, a.daily_at, a.budget_usd, ctx.actor)
    return {"schedule_id": sid}


def h_webhook(ctx, a: WebhookIn):
    validate_url(a.url)
    body = json.dumps({"text": a.text, "content": a.text}).encode()
    req = urllib.request.Request(a.url, data=body, headers={"Content-Type": "application/json", "User-Agent": "SASTRA"}, method="POST")
    with urllib.request.urlopen(req, timeout=20) as resp:  # noqa: S310 - validated + owner-approved
        status = resp.status
    _svc(ctx).privacy.record(a.url[:200], "webhook", "private", len(body), "approved webhook message", ctx.task_id)
    return {"status": status}


def h_project_delete(ctx, a: ProjectDeleteIn):
    _svc(ctx).projects.delete(a.project_id, ctx.actor)
    return {"deleted": a.project_id}


def register_builtin(reg: ToolRegistry) -> ToolRegistry:
    T = ToolSpec
    reg.register(T("calculator", "Exact arithmetic", CalcIn, h_calc))
    reg.register(T("memory_search", "Search governed memory (project-scoped)", MemSearchIn, h_mem_search, ["memory.read"], "private"))
    reg.register(T("memory_propose", "Propose a memory (saved as unvalidated/experimental)", MemProposeIn, h_mem_propose,
                   ["memory.write"], "private", side_effects=True, min_mode=Mode.ASSISTED))
    reg.register(T("web_search", "Search the public web via an authorised search API", SearchIn, h_search, ["network"], "public",
                   network=True, cost_usd=0.0, timeout_s=40))
    reg.register(T("fetch_page", "Fetch a public web page or PDF (SSRF-guarded, robots-respecting)", FetchIn, h_fetch,
                   ["network"], "public", network=True, timeout_s=40))
    reg.register(T("research_report", "Multi-source research with a cited report", ResearchIn, h_research, ["network"], "public",
                   network=True, timeout_s=300))
    reg.register(T("fs_list", "List files in the project workspace", FsListIn, h_fs_list, ["fs.read"]))
    reg.register(T("fs_read", "Read a file in the project workspace", FsReadIn, h_fs_read, ["fs.read"], "private"))
    reg.register(T("fs_write", "Write a file in the project workspace (auto-checkpointed, reversible)", FsWriteIn, h_fs_write,
                   ["fs.write"], side_effects=True, risk="medium", min_mode=Mode.ASSISTED))
    reg.register(T("code_run", "Run code in the isolated sandbox", CodeRunIn, h_code_run, ["exec"], side_effects=True,
                   risk="medium", min_mode=Mode.ASSISTED, timeout_s=660, approval=lambda a: a.network != "none"))
    reg.register(T("git_checkpoint", "Commit the workspace as a checkpoint", CheckpointIn, h_checkpoint, ["fs.write"],
                   side_effects=True, min_mode=Mode.ASSISTED))
    reg.register(T("git_log", "Show checkpoint history", FsListIn, h_git_log, ["fs.read"]))
    reg.register(T("git_rollback", "Restore the workspace to a checkpoint (current state is saved first)", RollbackIn, h_rollback,
                   ["fs.write"], side_effects=True, risk="medium", min_mode=Mode.ASSISTED))
    reg.register(T("secrets_scan", "Scan the workspace for exposed credentials", FsListIn, h_secrets, ["fs.read"]))
    reg.register(T("spreadsheet_analyze", "Profile a CSV/XLSX file: columns, types, stats, missing values, duplicates", SheetIn,
                   h_sheet, ["fs.read"], "private", timeout_s=120))
    reg.register(T("document_generate", "Write a Markdown/HTML/Word document into the project", DocIn, h_doc, ["fs.write"],
                   side_effects=True, min_mode=Mode.ASSISTED))
    reg.register(T("artifact_store", "Store a workspace file as a versioned, hashed artifact", ArtifactIn, h_artifact,
                   ["artifacts.write"], side_effects=True, min_mode=Mode.ASSISTED))
    reg.register(T("android_project_generate", "Generate a buildable Jetpack Compose project + APK CI workflow", AndroidIn,
                   h_android, ["fs.write"], side_effects=True, min_mode=Mode.ASSISTED))
    reg.register(T("notify", "Send an in-app notification to the owner", NotifyIn, h_notify, ["notify"], side_effects=True))
    reg.register(T("schedule_create", "Create a recurring background task", ScheduleIn, h_schedule, ["schedule"],
                   side_effects=True, risk="high", min_mode=Mode.ASSISTED))
    reg.register(T("webhook_send", "Send a message to an external webhook (Slack/Discord/ntfy…)", WebhookIn, h_webhook,
                   ["network", "external_send"], "private", network=True, reversible=False, risk="high",
                   side_effects=True, min_mode=Mode.ASSISTED))
    reg.register(T("project_delete", "Permanently delete a project, its files, artifacts and memory", ProjectDeleteIn,
                   h_project_delete, ["fs.delete"], "private", reversible=False, risk="critical", side_effects=True,
                   min_mode=Mode.ASSISTED))
    return reg
